"""Builds the test corpus: workbooks written by openpyxl, then the same
workbooks saved by LibreOffice Calc (headless), so the reader meets two
producers' ways of writing the same content. Run from this directory:

    python3 make.py

Needs openpyxl and `soffice` on the PATH. The files are generated here and
licensed as the repository is.
"""

import datetime
import os
import subprocess
import tempfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName


def budget():
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Q1", "Q2", "Total"])
    rows = [("Rent", 1200, 1200), ("Food", 431.5, 512.25), ("Travel", 0, 950)]
    for i, (name, a, b) in enumerate(rows, start=2):
        ws.append([name, a, b, f"=B{i}+C{i}"])
    ws.append(["Sum", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=4):
        for c in row:
            c.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 18
    ws.freeze_panes = "A2"
    ws["A2"].comment = Comment("Paid on the first", "Mehmet")
    ws.merge_cells("F1:G1")
    ws["F1"] = "Merged note"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ws.add_chart(chart, "F3")
    dates = wb.create_sheet("Dates")
    dates["A1"] = datetime.date(2026, 10, 3)
    dates["A2"] = datetime.datetime(2026, 10, 3, 14, 30)
    dates["A3"] = datetime.time(9, 15)
    dates["A4"] = 0.256
    dates["A4"].number_format = "0.0%"
    dates["A5"] = True
    dates["A6"] = "=1/0"
    dates["A7"] = "Türkçe ğüşıöç İ"
    dates["A8"] = "  spaced  "
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "secret"
    wb.defined_names["Rates"] = DefinedName("Rates", attr_text="Budget!$B$2:$B$4")
    wb.save("openpyxl-budget.xlsx")


def empty():
    wb = Workbook()
    wb.save("openpyxl-empty.xlsx")


def via_libreoffice(src, dst):
    with tempfile.TemporaryDirectory() as out:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, src],
            check=True,
            capture_output=True,
        )
        os.replace(os.path.join(out, os.path.basename(src)), dst)


if __name__ == "__main__":
    budget()
    empty()
    via_libreoffice("openpyxl-budget.xlsx", "libreoffice-budget.xlsx")
